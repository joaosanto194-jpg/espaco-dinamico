import os
import sys
import logging
import unicodedata
import re
from pathlib import Path
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import List, Tuple, Dict, Optional

# Bibliotecas de terceiros
from pypdf import PdfReader
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import PatternFill

# ----------------------------------------------------------------------
# Configuração de Logging Profissional
# ----------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] [%(name)s] %(message)s",
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler("operacao_conformidade.log", encoding="utf-8")
    ]
)
logger = logging.getLogger("ConformidadeEngine")

# ----------------------------------------------------------------------
# Configurações de Diretórios (Com suporte a variáveis de ambiente)
# ----------------------------------------------------------------------
PASTA_RAIZ = Path(os.getenv("PASTA_RAIZ", r"C:\Users\joaovitorsantos\Desktop\JOAO VITOR - LEITOS NO EDOCS"))
ARQUIVO_SAIDA = Path(os.getenv("ARQUIVO_SAIDA", r"C:\Users\joaovitorsantos\Desktop\pdleitos\resultado_conformidade.xlsx"))


def normalizar(texto: str) -> str:
    """Normaliza o texto removendo acentos e convertendo para minúsculas."""
    if not texto:
        return ""
    texto_nfd = unicodedata.normalize("NFD", texto)
    texto_sem_acento = "".join(c for c in texto_nfd if unicodedata.category(c) != "Mn")
    return texto_sem_acento.lower()


def extrair_texto(arquivo: Path) -> str:
    """Extrai texto de PDFs, DOCX, TXT ou CSV de forma segura com tratamento de exceções."""
    if not arquivo.exists():
        logger.warning(f"Ficheiro não encontrado no disco: {arquivo}")
        return ""
        
    extensao = arquivo.suffix.lower()
    logger.debug(f"Extraindo texto do ficheiro: {arquivo.name} ({extensao})")
    
    try:
        if extensao == ".pdf":
            leitor = PdfReader(str(arquivo))
            return "\n".join(
                pagina.extract_text() or ""
                for pagina in leitor.pages
            )

        elif extensao == ".docx":
            documento = Document(str(arquivo))
            return "\n".join(
                paragrafo.text for paragrafo in documento.paragraphs
            )

        elif extensao in [".txt", ".csv"]:
            return arquivo.read_text(encoding="utf-8", errors="ignore")

    except Exception as erro:
        logger.error(f"Erro crítico ao ler o ficheiro {arquivo.name}: {erro}")
        return f"ERRO: {erro}"

    return ""


def localizar_documentos(pasta: Path) -> Tuple[List[Path], List[Path]]:
    """Varre o diretório do processo categorizando autorizações e notas fiscais."""
    autorizacoes: List[Path] = []
    notas_fiscais: List[Path] = []

    try:
        for arquivo in pasta.rglob("*"):
            if not arquivo.is_file():
                continue

            nome = normalizar(arquivo.name)

            if "autorizacao" in nome or "aut" in nome:
                autorizacoes.append(arquivo)

            if "nota fiscal" in nome or "notafiscal" in nome or re.search(r"\bnf\b", nome) or "danfe" in nome:
                notas_fiscais.append(arquivo)
                
    except Exception as e:
        logger.error(f"Erro ao varrer a pasta {pasta.name}: {e}")

    return autorizacoes, notas_fiscais


def primeiro_grupo(texto: str, padroes: List[str]) -> Optional[str]:
    """Testa múltiplos padrões de regex e retorna o primeiro grupo capturado."""
    for padrao in padroes:
        resultado = re.search(padrao, texto, re.IGNORECASE | re.DOTALL)
        if resultado:
            return resultado.group(1).strip()
    return None


def extrair_dados(texto: str) -> Dict[str, Any]:
    """Extrai CNPJ, Número da NF e Valor utilizando expressões regulares robustas."""
    texto_normalizado = normalizar(texto)

    cnpj = primeiro_grupo(
        texto_normalizado,
        [
            r"cnpj\s*[:\-]?\s*(\d{2}\.?\d{3}\.?\d{3}/?\d{4}\-?\d{2})"
        ]
    )

    numero_nf = primeiro_grupo(
        texto_normalizado,
        [
            r"(?:numero|n[uú]mero|n[ºo]|nf)\s*[:\-]?\s*(\d{1,20})",
            r"nota fiscal\s*[:\-]?\s*(\d{1,20})"
        ]
    )

    valor = primeiro_grupo(
        texto_normalizado,
        [
            r"(?:valor total|valor da nota|valor autorizado|total)\s*[:\-]?\s*r?\$?\s*([\d\.,]+)"
        ]
    )

    return {
        "CNPJ": limpar_cnpj(cnpj),
        "Número NF": numero_nf,
        "Valor": converter_valor(valor)
    }


def limpar_cnpj(cnpj: Optional[str]) -> Optional[str]:
    """Remove formatação do CNPJ mantendo apenas os dígitos."""
    if not cnpj:
        return None
    somente_numeros = re.sub(r"\D", "", cnpj)
    return somente_numeros if len(somente_numeros) == 14 else (somente_numeros or None)


def converter_valor(valor: Optional[str]) -> Optional[Decimal]:
    """Converte string monetária brasileira para Decimal com segurança."""
    if not valor:
        return None

    try:
        # Remove pontos de milhar e troca vírgula decimal por ponto
        valor_limpo = valor.replace(".", "").replace(",", ".")
        return Decimal(valor_limpo)
    except (InvalidOperation, ValueError):
        return None


def comparar(autorizacao_txt: str, nota_fiscal_txt: str) -> List[str]:
    """Compara os dados extraídos e retorna a lista de pendências de conformidade."""
    pendencias = []

    if not autorizacao_txt:
        pendencias.append("Autorização não localizada")

    if not nota_fiscal_txt:
        pendencias.append("Nota fiscal não localizada")

    if not autorizacao_txt or not nota_fiscal_txt:
        return pendencias

    dados_autorizacao = extrair_dados(autorizacao_txt)
    dados_nf = extrair_dados(nota_fiscal_txt)

    if dados_autorizacao["CNPJ"] and dados_nf["CNPJ"] and dados_autorizacao["CNPJ"] != dados_nf["CNPJ"]:
        pendencias.append("CNPJ da autorização diferente da nota fiscal")

    if dados_autorizacao["Número NF"] and dados_nf["Número NF"] and dados_autorizacao["Número NF"] != dados_nf["Número NF"]:
        pendencias.append("Número da NF divergente")

    if dados_autorizacao["Valor"] is not None and dados_nf["Valor"] is not None and dados_autorizacao["Valor"] != dados_nf["Valor"]:
        pendencias.append(
            f"Valor divergente: autorização R$ {dados_autorizacao['Valor']} x NF R$ {dados_nf['Valor']}"
        )

    for campo in ["CNPJ", "Número NF", "Valor"]:
        if not dados_autorizacao[campo]:
            pendencias.append(f"{campo} ausente na autorização")
        if not dados_nf[campo]:
            pendencias.append(f"{campo} ausente na nota fiscal")

    return pendencias


def analisar_processo(pasta: Path) -> Dict[str, str]:
    """Orquestra a análise individual de um processo/pasta."""
    logger.info(f"Analisando processo: {pasta.name}")
    autorizacoes, notas_fiscais = localizar_documentos(pasta)

    autorizacao = autorizacoes[0] if autorizacoes else None
    nota_fiscal = notas_fiscais[0] if notas_fiscais else None

    texto_autorizacao = extrair_texto(autorizacao) if autorizacao else ""
    texto_nota = extrair_texto(nota_fiscal) if nota_fiscal else ""

    pendencias = comparar(texto_autorizacao, texto_nota)

    return {
        "Processo": pasta.name,
        "Autorização": autorizacao.name if autorizacao else "",
        "Nota Fiscal": nota_fiscal.name if nota_fiscal else "",
        "Status": "Em condição de pagamento" if not pendencias else "Não apto para pagamento",
        "Pendências": " | ".join(pendencias) or "Autorização de acordo com a NF"
    }


def gerar_planilha() -> None:
    """Executa o fluxo completo do motor de conformidade e exporta para Excel."""
    if not PASTA_RAIZ.exists():
        logger.critical(f"Pasta raiz não encontrada: {PASTA_RAIZ}")
        raise FileNotFoundError(f"Pasta não encontrada: {PASTA_RAIZ}")

    resultados = []
    processos = [item for item in PASTA_RAIZ.iterdir() if item.is_dir()]

    logger.info(f"Total de processos encontrados para análise: {len(processos)}")

    for processo in processos:
        resultados.append(analisar_processo(processo))

    planilha = Workbook()
    aba = planilha.active
    aba.title = "Conformidade"

    colunas = ["Processo", "Autorização", "Nota Fiscal", "Status", "Pendências"]
    aba.append(colunas)

    for resultado in resultados:
        aba.append([resultado[coluna] for coluna in colunas])

    # Estilização condicional das células de status
    fill_verde = PatternFill(fill_type="solid", fgColor="C6EFCE")
    fill_vermelho = PatternFill(fill_type="solid", fgColor="FFC7CE")

    for linha in range(2, aba.max_row + 1):
        status = aba.cell(linha, 4)
        if status.value == "Em condição de pagamento":
            status.fill = fill_verde
        else:
            status.fill = fill_vermelho

    # Ajuste automático de largura das colunas
    for coluna in aba.columns:
        largura = max(len(str(celula.value or "")) for celula in coluna) + 2
        aba.column_dimensions[coluna[0].column_letter].width = min(largura, 80)

    # Tratamento de salvamento seguro (evita crash se o Excel estiver aberto)
    try:
        ARQUIVO_SAIDA.parent.mkdir(parents=True, exist_ok=True)
        planilha.save(ARQUIVO_SAIDA)
        logger.info(f"Relatório gerado com sucesso em: {ARQUIVO_SAIDA}")
    except PermissionError:
        logger.error(f"Erro de permissão: O ficheiro {ARQUIVO_SAIDA} está aberto no Excel. Feche-o e tente novamente.")
    except Exception as e:
        logger.error(f"Erro ao salvar o ficheiro Excel: {e}")


if __name__ == "__main__":
    gerar_planilha()